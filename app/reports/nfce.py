"""Excel dentro da pasta de cada empresa, seguindo o relatório da automação NFSe."""
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.automation.sat_nfce_xml import preparar_xml
from app.output_layout import nomes_saida, pastas_empresa

NS = {"n": "http://www.portalfiscal.inf.br/nfe"}
BLUE, TEAL, PALE, WHITE = "17324D", "087F8C", "EAF3F8", "FFFFFF"
MONEY = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
CANCEL_TYPES = {"110111", "110112"}


def parse(data):
    root = etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False))
    if root.getroottree().docinfo.doctype:
        raise ValueError("DOCTYPE não permitido")
    return root


def txt(root, path):
    node = root.find(path, NS) if root is not None else None
    return (node.text or "").strip() if node is not None else ""


def number(root, path, required=False):
    text = txt(root, path)
    if not text and not required:
        return Decimal(0)
    try:
        value = Decimal(text)
        if not value.is_finite():
            raise InvalidOperation()
        return value
    except InvalidOperation:
        raise ValueError("Valor numérico inválido no XML") from None


def cell(ws, row, col, value):
    c = ws.cell(row, col, value)
    if isinstance(value, str):
        c.data_type = "s"  # Nomes/produtos/chaves são texto, nunca fórmulas.
    return c


def base_sheet(ws, title, subtitle, columns):
    width = len(columns)
    ws.sheet_view.showGridLines = False
    for row, text in ((1, title), (2, subtitle)):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        c = cell(ws, row, 1, text)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.font = Font(name="Calibri", size=16 if row == 1 else 10, bold=row == 1, color=WHITE)
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 32 if row == 1 else 24
    for i, (label, size, _key, _kind) in enumerate(columns, 1):
        c = cell(ws, 4, i, label)
        c.fill = PatternFill("solid", fgColor=TEAL)
        c.font = Font(name="Calibri", bold=True, color=WHITE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = size
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "C5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    ws.oddFooter.center.text = "ZNX • NFC-e | Página &P de &N"


def table(ws, title, subtitle, columns, rows, total=True):
    base_sheet(ws, title, subtitle, columns)
    sums = defaultdict(Decimal)
    for r, item in enumerate(rows, 5):
        for i, (_label, _size, key, kind) in enumerate(columns, 1):
            value = item.get(key, "")
            c = cell(ws, r, i, value)
            c.font = Font(name="Calibri", size=10)
            c.alignment = Alignment(vertical="top")
            if r % 2:
                c.fill = PatternFill("solid", fgColor=PALE)
            if kind == "money":
                c.number_format = MONEY
                if isinstance(value, (Decimal, int, float)):
                    sums[i] += Decimal(str(value))
            elif kind == "decimal":
                c.number_format = "0.0000"
            elif kind == "text":
                c.number_format = "@"
            if key == "situacao" and value == "Cancelada":
                c.font = Font(name="Calibri", bold=True, color="B42318")
    ws.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{max(4, len(rows)+4)}"
    if total:
        r = len(rows) + 6
        for i in range(1, len(columns)+1):
            c = cell(ws, r, i, f"TOTAL · {len(rows)} registros" if i == 1 else sums.get(i, ""))
            c.fill = PatternFill("solid", fgColor=BLUE)
            c.font = Font(name="Calibri", bold=True, color=WHITE)
            if i in sums:
                c.number_format = MONEY


NOTE_COLUMNS = [
    ("Número", 12, "numero", "text"), ("Série", 8, "serie", "text"),
    ("Emissão", 26, "emissao", "text"), ("Situação observada", 32, "situacao", "text"),
    ("Valor da NFC-e", 19, "valor", "money"), ("Produtos", 18, "produtos", "money"),
    ("Desconto", 16, "desconto", "money"), ("Frete", 16, "frete", "money"),
    ("Outras despesas", 18, "outros", "money"), ("Base ICMS", 17, "base_icms", "money"),
    ("ICMS", 16, "icms", "money"), ("ICMS ST", 16, "st", "money"),
    ("PIS", 16, "pis", "money"), ("COFINS", 16, "cofins", "money"),
    ("Tributos aproximados", 21, "tributos", "money"),
    ("Emitente CNPJ", 20, "emitente_doc", "text"), ("Emitente", 38, "emitente", "text"),
    ("Destinatário CPF/CNPJ", 23, "dest_doc", "text"), ("Destinatário", 32, "dest_nome", "text"),
    ("Código do protocolo", 20, "cstat", "text"), ("Protocolo", 23, "protocolo", "text"),
    ("Chave de acesso", 49, "chave", "text"), ("Arquivo XML", 42, "arquivo", "text"),
]
ITEM_COLUMNS = [
    ("Número", 12, "numero", "text"), ("Série", 8, "serie", "text"),
    ("Situação observada", 32, "situacao", "text"), ("Item", 8, "item", "text"),
    ("Código produto", 18, "codigo", "text"), ("Descrição", 48, "descricao", "text"),
    ("NCM", 12, "ncm", "text"), ("CFOP", 10, "cfop", "text"), ("Unidade", 10, "unidade", "text"),
    ("Quantidade", 16, "quantidade", "decimal"), ("Valor unitário", 18, "unitario", "decimal"),
    ("Valor produto", 18, "valor", "money"), ("Desconto", 16, "desconto", "money"),
    ("CST / CSOSN ICMS", 18, "cst", "text"), ("ICMS", 16, "icms", "money"),
    ("PIS", 16, "pis", "money"), ("COFINS", 16, "cofins", "money"),
    ("Chave de acesso", 49, "chave", "text"),
]
EVENT_COLUMNS = [
    ("Data do evento", 27, "data", "text"), ("Tipo", 12, "tipo", "text"),
    ("Descrição", 34, "descricao", "text"), ("Código de retorno", 18, "cstat", "text"),
    ("Dentro do período", 19, "no_periodo", "text"), ("Nota neste relatório", 21, "nota_presente", "text"),
    ("Protocolo", 23, "protocolo", "text"), ("Chave de acesso", 49, "chave", "text"),
    ("NSU", 17, "nsu", "text"),
]


def read_client(folder, cnpj, inicio, fim):
    notes, items, seen = [], [], set()
    for path in sorted(folder.glob("NFCe_*.xml")):
        root = parse(preparar_xml(parse(path.read_bytes())))
        inf = root.find("n:NFe/n:infNFe", NS)
        if inf is None or txt(inf, "n:ide/n:mod") != "65":
            raise ValueError("XML não é NFC-e")
        key = inf.get("Id", "").removeprefix("NFe")
        if not key or key in seen:
            raise ValueError("Chave ausente ou duplicada")
        seen.add(key)
        emit_doc = txt(inf, "n:emit/n:CNPJ")
        issued = txt(inf, "n:ide/n:dhEmi")
        if emit_doc != cnpj:
            raise ValueError("XML pertence a outra empresa")
        if not inicio <= issued[:10] <= fim:
            continue
        totals = inf.find("n:total/n:ICMSTot", NS)
        note = dict(numero=txt(inf,"n:ide/n:nNF"), serie=txt(inf,"n:ide/n:serie"), emissao=issued,
            chave=key, arquivo=path.name, valor=number(totals,"n:vNF",True),
            emitente_doc=emit_doc, emitente=txt(inf,"n:emit/n:xNome"),
            dest_doc=txt(inf,"n:dest/n:CNPJ") or txt(inf,"n:dest/n:CPF"), dest_nome=txt(inf,"n:dest/n:xNome"),
            cstat=txt(root,"n:protNFe/n:infProt/n:cStat"), protocolo=txt(root,"n:protNFe/n:infProt/n:nProt"))
        for name, tag in {"produtos":"vProd","desconto":"vDesc","frete":"vFrete","outros":"vOutro",
                          "base_icms":"vBC","icms":"vICMS","st":"vST","pis":"vPIS",
                          "cofins":"vCOFINS","tributos":"vTotTrib"}.items():
            note[name] = number(totals, "n:"+tag)
        notes.append(note)
        for det in inf.findall("n:det", NS):
            prod, tax = det.find("n:prod", NS), det.find("n:imposto", NS)
            item = dict(numero=note["numero"], serie=note["serie"], chave=key, item=det.get("nItem", ""))
            for name, tag in {"codigo":"cProd","descricao":"xProd","ncm":"NCM","cfop":"CFOP","unidade":"uCom"}.items():
                item[name] = txt(prod,"n:"+tag)
            for name, tag in {"quantidade":"qCom","unitario":"vUnCom","valor":"vProd","desconto":"vDesc"}.items():
                item[name] = number(prod,"n:"+tag)
            item["cst"] = txt(tax,"n:ICMS/*/n:CST") or txt(tax,"n:ICMS/*/n:CSOSN")
            for name, tag in {"icms":"ICMS/*/n:vICMS","pis":"PIS/*/n:vPIS","cofins":"COFINS/*/n:vCOFINS"}.items():
                item[name] = number(tax,"n:"+tag)
            items.append(item)
    events, event_ids = [], set()
    # Inclui eventos posteriores vinculados às notas do período, se já distribuídos.
    sources = list(folder.glob("Evento_*.xml"))
    for path in sorted(sources):
        root = parse(path.read_bytes())
        docs = [(root, path.stem.rsplit("_", 1)[-1])]
        for doc, nsu in docs:
            inf = doc.find("n:evento/n:infEvento", NS)
            if inf is None:
                continue
            key, day = txt(inf,"n:chNFe"), txt(inf,"n:dhEvento")
            if key not in seen and not inicio <= day[:10] <= fim:
                continue
            parse(preparar_xml(doc))
            identity = inf.get("Id", "")
            if not identity:
                raise ValueError("Evento sem identificação")
            if identity in event_ids:
                continue
            event_ids.add(identity)
            returned = doc.find("n:retEvento/n:infEvento",NS)
            matches = txt(returned,"n:chNFe") == key and txt(returned,"n:tpEvento") == txt(inf,"n:tpEvento")
            events.append(dict(data=day, tipo=txt(inf,"n:tpEvento"), descricao=txt(inf,"n:detEvento/n:descEvento"),
                cstat=txt(returned,"n:cStat"), protocolo=txt(returned,"n:nProt"), chave=key, nsu=nsu,
                no_periodo="Sim" if inicio <= day[:10] <= fim else "Não", nota_presente="Sim" if key in seen else "Não",
                cancelado=matches and txt(inf,"n:tpEvento") in CANCEL_TYPES and txt(returned,"n:cStat") in {"135","155"}))
    cancelled = {e["chave"] for e in events if e["cancelado"]}
    for obj in notes + items:
        obj["situacao"] = "Cancelada" if obj["chave"] in cancelled else "Sem cancelamento identificado"
    return notes, items, events


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name("."+uuid4().hex+".xlsx")
    try:
        wb.save(temp)
        temp.replace(path)
    finally:
        wb.close()
        temp.unlink(missing_ok=True)


def gerar_cliente(folder, nome, cnpj, inicio, fim, resultado, destino):
    notes, items, events = read_client(Path(folder), cnpj, inicio, fim)
    if resultado.get("status") != "erro" and len(notes) != resultado.get("notas", len(notes)):
        raise ValueError("Contagem do relatório diverge do download")
    cancelled = [n for n in notes if n["situacao"] == "Cancelada"]
    gross = sum((n["valor"] for n in notes), Decimal(0))
    canc_value = sum((n["valor"] for n in cancelled), Decimal(0))
    summary = dict(notas=len(notes), itens=len(items), eventos_periodo=sum(e["no_periodo"] == "Sim" for e in events),
                   eventos_relacionados=len(events), canceladas=len(cancelled), valor=str(gross),
                   valor_canceladas=str(canc_value), valor_sem_cancelamento=str(gross-canc_value))
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    subtitle = f"{nome} | CNPJ {cnpj} | {inicio} a {fim}"
    base_sheet(ws, "ZNX • Conferência NFC-e", subtitle,
               [("Indicador",40,"", ""),("Resultado",65,"", "")])
    rows = [
        ("Resultado da execução", resultado.get("status", "")),
        ("Consulta / download", resultado.get("message", "")),
        ("NFC-e do período", len(notes)), ("Itens",len(items)),
        ("Eventos ocorridos no período", summary["eventos_periodo"]),
        ("Eventos incluídos (período + relacionados)",len(events)),
        ("Cancelamentos identificados nas notas",len(cancelled)),
        ("Valor total das NFC-e",gross), ("Valor das canceladas identificadas",canc_value),
        ("Valor sem cancelamento identificado",gross-canc_value),
        ("Último NSU distribuído",str(resultado.get("ultimo_nsu", ""))),
        ("Origem desta execução", "Arquivos locais" if resultado.get("origem")=="cache" else "API SEF"),
        ("Próxima consulta permitida",resultado.get("proxima_consulta", "")),
        ("Gerado em",datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Conferência", "Digest e assinatura dos XMLs conferidos. Valores transcritos dos XMLs."),
        ("Cobertura", "Documentos disponíveis na distribuição. Não comprova ausência de notas fora da janela da SEF."),
        ("Situação das notas", "Cancelamento identificado por evento aceito (135/155). Eventos relacionados posteriores são considerados quando presentes no cache."),
        ("Valores e tributos", "Totais incluem notas canceladas na aba NFC-e; a aba Canceladas é um recorte, não deve ser somada novamente."),
        ("Limite da conferência", "Sem cancelamento identificado não confirma situação fiscal atual. Não é apuração de impostos nem validação da cadeia ICP-Brasil."),
    ]
    for r, (label,value) in enumerate(rows,5):
        cell(ws,r,1,label).font = Font(name="Calibri",bold=True,color=BLUE)
        c=cell(ws,r,2,value)
        c.alignment=Alignment(wrap_text=True,vertical="center")
        if isinstance(value,Decimal): c.number_format=MONEY
        if r % 2:
            for col in (1,2): ws.cell(r,col).fill=PatternFill("solid",fgColor=PALE)
        ws.row_dimensions[r].height=46 if isinstance(value,str) and len(value)>65 else 25
    for title, columns, data in [("NFC-e",NOTE_COLUMNS,notes),("Canceladas",NOTE_COLUMNS,cancelled),
                                  ("Itens",ITEM_COLUMNS,items),("Eventos",EVENT_COLUMNS,events)]:
        table(wb.create_sheet(title), "ZNX • "+title, subtitle, columns, data)
    save(wb, Path(destino))
    return summary


def gerar_execucao(params, results, pasta_saida, job_id, status):
    """Cria somente Empresa/Relatorio e Empresa/XMLs para cada cliente."""
    mapped = {r["cnpj"]:r for r in results}
    output_names = nomes_saida(params.clientes)
    files, folders, errors = {}, {}, []
    for client in params.clientes:
        cnpj, name = client["documento"], client["nome"]
        result = mapped.get(cnpj, {})
        root, period_folder, report_folder, xml_folder = pastas_empresa(
            pasta_saida, output_names[cnpj], params.data_inicio, params.data_fim)
        report_folder.mkdir(parents=True, exist_ok=True)
        xml_folder.mkdir(parents=True, exist_ok=True)
        effective = {"cnpj": cnpj, "nome": name, "status": "não processado",
                     "message": "Execução interrompida antes deste cliente", "notas": 0, **result}
        try:
            path = report_folder / f"Relatorio_NFCe_{params.data_inicio}_a_{params.data_fim}.xlsx"
            source_folder = Path(result.get("pasta") or xml_folder)
            gerar_cliente(source_folder, name, cnpj, params.data_inicio, params.data_fim, effective, path)
            files[cnpj] = str(path)
        except Exception:
            errors.append(cnpj)
        folders[cnpj] = {"empresa": str(root), "periodo": str(period_folder),
                         "relatorio": str(report_folder), "xml": str(xml_folder)}
    return {"clientes": files, "pastas": folders, "erros": errors}
