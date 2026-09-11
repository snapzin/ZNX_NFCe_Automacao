"""
FastAPI backend para NFC-e SAT Automacao.
Expõe endpoints REST para o frontend Electron.
"""
import hmac
import importlib
import logging
import os
import re
import sys
import threading
import uuid
import time as _time
from collections import deque
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, HTTPException, Request
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
import uvicorn

# =============================================================================
# Logging — PRIMEIRO, antes de qualquer import que possa usar logger
# =============================================================================
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Suporte a bundle PyInstaller: config.py e módulos ficam ao lado do server.exe
if getattr(sys, "frozen", False):
    _app_dir = str(Path(sys.executable).parent)
    if _app_dir not in sys.path:
        sys.path.insert(0, _app_dir)
else:
    sys.path.insert(0, str(Path(__file__).parent.parent))

try:
    from app.config.config_defaults import apply_defaults, default_values, ensure_config_file, project_root
except ImportError as e:
    logger.warning(f"Falha ao importar config_defaults: {e}")
    apply_defaults = None
    default_values = lambda: {}
    ensure_config_file = lambda: Path(__file__).parent.parent / "config.py"

try:
    _config_path = ensure_config_file()
    sys.path.insert(0, str(_config_path.parent))
except Exception as e:
    logger.warning("Falha ao criar config.py padrao: %s", e)

try:
    import config
    if apply_defaults:
        apply_defaults(config)
except ImportError:
    logger.warning("Falha ao importar config")
    class config:
        pass
    if apply_defaults:
        apply_defaults(config, create_dirs=False)

from app.automation.sat_automacao import ExecucaoCancelada, executar_local, preparar_parametros
from app.automation.sat_session import CredenciaisSAT


# =============================================================================
# Job management
# =============================================================================
@dataclass
class Job:
    id: str
    status: str  # "running", "ok", "cancelado", "erro"
    logs: deque = None
    thread: Optional[threading.Thread] = None
    cancel_event: Optional[threading.Event] = None
    resultado: Optional[dict] = None
    progress: dict = None
    authentication: Optional[dict] = None
    created_at: float = 0.0

    def __post_init__(self):
        if self.progress is None:
            self.progress = {"total": 0, "processed": 0, "clients": [], "stage": "Preparando"}
        if self.logs is None:
            self.logs = deque(maxlen=2000)
        if self.cancel_event is None:
            self.cancel_event = threading.Event()
        if self.created_at == 0.0:
            self.created_at = _time.time()

JOBS = {}
_MAX_JOBS = 50
_JOB_TTL_S = 7200  # 2 horas


def _prune_jobs() -> None:
    """Remove jobs finalizados mais antigos quando o limite é atingido."""
    if len(JOBS) <= _MAX_JOBS:
        return
    cutoff = _time.time() - _JOB_TTL_S
    to_remove = [
        jid for jid, j in JOBS.items()
        if j.status != "running" and j.created_at < cutoff
    ]
    for jid in to_remove:
        del JOBS[jid]
    # Se ainda acima do limite, remove os mais antigos independente do TTL
    if len(JOBS) > _MAX_JOBS:
        finished = sorted(
            [(jid, j) for jid, j in JOBS.items() if j.status != "running"],
            key=lambda x: x[1].created_at,
        )
        for jid, _ in finished[:len(JOBS) - _MAX_JOBS]:
            del JOBS[jid]

class JobLogHandler(logging.Handler):
    """Handler que adiciona logs a um Job específico."""
    def __init__(self, job_id: str):
        super().__init__()
        self.job_id = job_id
        self.setFormatter(
            logging.Formatter(
                fmt="%(asctime)s  %(levelname)-8s  %(name)s  %(message)s",
                datefmt="%H:%M:%S",
            )
        )

    def emit(self, record: logging.LogRecord) -> None:
        if (self.job_id in JOBS and record.name == "sat"
                and JOBS[self.job_id].thread
                and record.thread == JOBS[self.job_id].thread.ident):
            msg = self.format(record)
            JOBS[self.job_id].logs.append({
                "timestamp": record.created,
                "level": record.levelname,
                "message": msg,
            })

# =============================================================================
# FastAPI app
# =============================================================================
app = FastAPI(title="NFC-e SAT Automacao API")


@app.exception_handler(RequestValidationError)
async def invalid_request(request, exc):
    # O erro padrão pode incluir o corpo recebido, inclusive campos secretos.
    return JSONResponse(status_code=422, content={"detail": "Formato de requisição inválido"})

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://127.0.0.1:5173", "http://localhost:5173", "app://"],
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type", "x-api-token"],
)

# Token gerado pelo Electron e passado via env var NFSE_API_TOKEN.
# Se não definido (modo dev sem Electron), a auth é desativada.
_API_TOKEN = os.environ.get("NFSE_API_TOKEN", "")

@app.middleware("http")
async def _require_token(request: Request, call_next):
    if request.url.path == "/health":
        return await call_next(request)
    if _API_TOKEN:
        token = request.headers.get("x-api-token", "")
        if not token or not hmac.compare_digest(token.encode(), _API_TOKEN.encode()):
            return JSONResponse(status_code=401, content={"error": "Não autorizado"})

    return await call_next(request)

# =============================================================================
# Endpoints
# =============================================================================

@app.get("/health")
async def health():
    """Health check — Electron aguarda isso para saber que o backend está pronto."""
    return {"ok": True}


@app.get("/health/auth")
async def authenticated_health():
    """Confirma que Electron e backend compartilham o mesmo token de sessao."""
    return {"ok": True, "authenticated": True}

@app.get("/config")
async def get_config():
    """Retorna seções + valores atuais de config.py."""
    sections = [
        {"title": "Caminhos locais", "fields": [
            ("XLSX_PATH", "Planilha de clientes", "path"),
            ("PASTA_SAIDA", "Pasta de saída", "dir"),
        ]},
        {"title": "SAT / HTTP sem navegador", "fields": [
            ("SAT_HTTP_TIMEOUT_S", "Tempo máximo da resposta HTTP (s)", "int"),
        ]},
        {"title": "Seletores SAT — alterar após observar o portal", "fields": [
            (key, key.removeprefix("SAT_SELECTOR_").replace("_", " ").capitalize(), "text")
            for key in default_values() if key.startswith("SAT_SELECTOR_")
        ]},
    ]

    values = {}
    for section in sections:
        for key, label, tipo in section["fields"]:
            val = getattr(config, key, "")
            values[key] = "••••••••" if tipo == "secret" else str(val)

    return {"sections": sections, "values": values}

@app.post("/config")
async def set_config(body: dict):
    """Salva valores em config.py (com backup .bak)."""
    config_path = ensure_config_file()
    if not config_path.exists():
        try:
            config_path = ensure_config_file(config_path.parent)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Falha ao criar config.py: {e}")
    if not config_path.exists():
        raise HTTPException(status_code=404, detail="config.py não encontrado")

    try:
        original = config_path.read_text(encoding="utf-8")
        novo = original
        defaults = default_values()

        for chave, valor in body.items():
            if chave == "SAT_LOGIN_URL" or chave not in defaults:
                continue
            if not hasattr(config, chave):
                continue
            if str(valor).strip() == "" and defaults.get(chave) not in ("", None):
                valor = defaults[chave]
            atual = getattr(config, chave)
            if isinstance(atual, bool):
                literal = "True" if str(valor).lower() in ("true", "1", "sim") else "False"
            elif isinstance(atual, int):
                literal = str(int(valor))
            else:
                literal = repr(str(valor))

            pattern = re.compile(
                rf"^(?P<prefix>{re.escape(chave)}\s*=\s*)(?P<value>.+?)(?P<suffix>\s*(?:#.*)?)$",
                re.MULTILINE,
            )
            novo, count = pattern.subn(
                lambda m: f"{m.group('prefix')}{literal}{m.group('suffix')}",
                novo,
                count=1,
            )
            if count == 0:
                novo = f"{novo.rstrip()}\n{chave} = {literal}\n"

        backup = config_path.with_name(f"config.py.bak.{int(_time.time())}")
        backup.write_text(original, encoding="utf-8")

        # Remove backups mais antigos (mantém últimos 5)
        old_backups = sorted(config_path.parent.glob("config.py.bak.*"),
                             key=lambda p: p.stat().st_mtime)
        for old in old_backups[:-5]:
            try:
                old.unlink()
            except Exception:
                pass
        compile(novo, str(config_path), "exec")
        temporary = config_path.with_suffix(".tmp")
        temporary.write_text(novo, encoding="utf-8")
        temporary.replace(config_path)

        importlib.reload(config)
        if apply_defaults:
            apply_defaults(config)

        return {"ok": True, "backupPath": str(backup)}
    except Exception as e:
        logger.error("Falha ao salvar configuração")
        raise HTTPException(status_code=500, detail="Falha ao salvar configuração")

@app.get("/paths/status")
async def paths_status():
    return {"needsSetup": False, "suggestions": {"PASTA_SAIDA": config.PASTA_SAIDA},
            "saida": {"path": config.PASTA_SAIDA, "ok": Path(config.PASTA_SAIDA).is_dir()}}


def _resolver_xlsx_path() -> Path:
    """Resolve XLSX_PATH relativo ao diretorio do config.py."""
    raw = str(getattr(config, "XLSX_PATH", "clientes.xlsx")).strip() or "clientes.xlsx"
    p = Path(raw)
    if p.is_absolute():
        return p
    config_dir = Path(getattr(config, "__file__", __file__)).parent
    return config_dir / raw


def _normalizar_clientes(clientes) -> list[dict]:
    out = []
    for c in clientes or []:
        if not isinstance(c, dict):
            continue
        doc = re.sub(r"\D", "", str(c.get("documento", c.get("cnpj", c.get("cpf", "")))).strip())[:32]
        nome = str(c.get("nome", "")).strip()[:200]
        # 'ativo' controla quais clientes entram na execucao (checkbox na tela).
        # Ausente = True (retrocompatibilidade: tudo ativo por padrao).
        ativo = c.get("ativo", True)
        if isinstance(ativo, str):
            ativo = ativo.strip().lower() not in ("nao", "não", "false", "0", "n", "")
        if doc or nome:
            out.append({"documento": doc, "nome": nome, "ativo": bool(ativo)})
    return out


def _ler_clientes_local() -> dict:
    xlsx_path = _resolver_xlsx_path()
    clientes = _parse_xlsx_clientes(xlsx_path) if xlsx_path.exists() else []
    return {"clientes": clientes, "path": str(xlsx_path), "source": "local"}


def _salvar_clientes_local(clientes: list[dict]) -> dict:
    from openpyxl import Workbook

    xlsx_path = _resolver_xlsx_path()
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append(["NOME", "CNPJ"])
    for c in clientes:
        ws.append([c.get("nome", ""), c.get("documento", "")])
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    temporary = xlsx_path.with_name(xlsx_path.stem + ".tmp.xlsx")
    wb.save(temporary)
    temporary.replace(xlsx_path)
    wb.close()

    return {"ok": True, "salvos": len(clientes), "path": str(xlsx_path), "source": "local"}


# ── Overlay LOCAL da selecao (ativo/inativo) por CNPJ ────────────────────────
# Seleção local preservada no formato existente, por CPF/CNPJ.
def _overlay_path() -> Path:
    return _resolver_xlsx_path().with_name("clientes_selecao.json")


def _ler_overlay() -> dict:
    import json
    p = _overlay_path()
    if not p.exists():
        return {}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
        return {re.sub(r"\D", "", str(k)): bool(v) for k, v in (data or {}).items()}
    except Exception as e:  # noqa: BLE001
        logger.warning("Falha ao ler selecao de clientes: %s", e)
        return {}


def _salvar_overlay(clientes: list[dict]) -> None:
    import json
    sel = {}
    for c in clientes or []:
        doc = re.sub(r"\D", "", str(c.get("documento", "")))
        if doc:
            sel[doc] = bool(c.get("ativo", True))
    try:
        _overlay_path().write_text(json.dumps(sel), encoding="utf-8")
    except OSError as e:
        logger.warning("Falha ao salvar selecao de clientes: %s", e)


def _aplicar_overlay(clientes: list[dict]) -> list[dict]:
    """Sobrescreve 'ativo' de cada cliente com a selecao local salva (por CNPJ).
    Clientes sem entrada no overlay mantem o proprio 'ativo' (default True)."""
    ov = _ler_overlay()
    if not ov:
        return clientes
    for c in clientes:
        doc = re.sub(r"\D", "", str(c.get("documento", "")))
        if doc in ov:
            c["ativo"] = ov[doc]
    return clientes


def _aplicar_dados_locais(clientes: list[dict]) -> list[dict]:
    return _aplicar_overlay(_normalizar_clientes(clientes))


def _clientes_atuais() -> list[dict]:
    """Carrega os clientes locais selecionados para a execução."""
    try:
        return _aplicar_dados_locais(_ler_clientes_local().get("clientes") or [])
    except Exception as local_err:
        logger.warning("Falha ao ler clientes locais p/ execucao: %s", local_err)
        return []


@app.get("/clientes")
async def listar_clientes(request: Request):
    """Lê o cadastro local com seleção por empresa."""
    try:
        local = _ler_clientes_local()
        local["clientes"] = _aplicar_dados_locais(local.get("clientes") or [])
        if not local.get("location"):
            local["location"] = f"Arquivo local: {local.get('path')}"
        return local
    except Exception as e:
        logger.exception("Falha ao listar clientes")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/clientes")
async def salvar_clientes(request: Request, payload: dict):
    """Grava o cadastro local, sem dependência de servidor de licença."""
    try:
        clientes = _normalizar_clientes(payload.get("clientes") or [])
        docs = set()
        for cliente in clientes:
            doc = cliente["documento"]
            if len(doc) != 14 or not cliente["nome"]:
                raise HTTPException(400, "Informe CNPJ com 14 dígitos e nome de cada cliente")
            if doc in docs:
                raise HTTPException(400, "CNPJ duplicado")
            docs.add(doc)
        _salvar_overlay(clientes)

        return _salvar_clientes_local(clientes)
    except Exception as e:
        if isinstance(e, HTTPException):
            raise
        logger.error("Falha ao salvar clientes")
        raise HTTPException(status_code=500, detail="Falha ao salvar clientes")


def _parse_xlsx_clientes(path: Path) -> list[dict]:
    """Importa NOME/CNPJ em qualquer ordem. Colunas antigas de acesso são ignoradas."""
    import unicodedata
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    def header(value):
        return "".join(c for c in unicodedata.normalize("NFD", str(value or "").strip().lower())
                       if unicodedata.category(c) != "Mn")
    headers = [header(c) for c in rows[0]]
    idx_doc = next((i for i, h in enumerate(headers) if h in ("cnpj", "documento")), None)
    idx_nome = next((i for i, h in enumerate(headers) if h in ("nome", "nome da empresa", "razao social", "cliente")), None)
    if idx_doc is None or idx_nome is None:
        raise ValueError("A planilha precisa das colunas NOME e CNPJ")
    idx_ativo = next((i for i, h in enumerate(headers) if h == "ativo"), None)
    clients = []
    for row in rows[1:]:
        if not any(v is not None and str(v).strip() for v in row):
            continue
        def cell(index, default=""):
            return row[index] if index is not None and index < len(row) and row[index] is not None else default
        doc = cell(idx_doc)
        if isinstance(doc, int):
            doc = str(doc).zfill(14)
        clients.append({"documento": str(doc), "nome": str(cell(idx_nome)), "ativo": cell(idx_ativo, True)})
    normalized = _normalizar_clientes(clients)
    if any(len(c["documento"]) != 14 or not c["nome"] for c in normalized):
        raise ValueError("Preencha nome e CNPJ de cada cliente")
    return normalized


@app.post("/clientes/importar")
async def importar_clientes(payload: dict):
    """Le um XLSX (CNPJ/CPF + Nome) do caminho informado e DEVOLVE os clientes
    (nao salva — o frontend mescla com a lista e o usuario clica em Salvar)."""
    caminho = str(payload.get("path") or "").strip()
    if not caminho:
        raise HTTPException(status_code=400, detail="Caminho do arquivo nao informado.")
    p = Path(caminho)
    if not p.exists():
        raise HTTPException(status_code=404, detail=f"Arquivo nao encontrado: {caminho}")
    try:
        clientes = _parse_xlsx_clientes(p)
    except Exception as e:
        logger.error("Falha ao importar planilha")
        raise HTTPException(status_code=400, detail="Não foi possível ler a planilha XLSX/XLSM")
    return {"ok": True, "clientes": clientes, "total": len(clientes)}


def _filtrar_clientes_por_login(todos_clientes: list[dict], tipo_login: str) -> list[dict]:
    if tipo_login != "usuario_senha":
        raise HTTPException(400, "O acesso ao SAT usa usuário e senha do contador")
    selected = [c for c in todos_clientes if c.get("ativo", True)]
    if not selected:
        raise HTTPException(400, "Nenhum cliente ativo selecionado na aba Clientes")
    return selected


@app.post("/executar")
async def start_execution(request: Request, body: dict):
    """Uma autenticação do contador; clientes são os CNPJs alvo da consulta."""
    if any(j.status == "running" for j in JOBS.values()):
        raise HTTPException(409, "Já existe uma execução em andamento")
    if body.get("modo", "consulta") != "consulta":
        raise HTTPException(400, "Use o fluxo de consulta com usuário e senha")
    _prune_jobs()
    clients = _filtrar_clientes_por_login(
        _clientes_atuais(),
        body.get("tipoLogin", "usuario_senha"),
    )
    try:
        params = preparar_parametros(body.get("dataInicio"), body.get("dataFim"),
                                    body.get("cnpjs"), clientes=clients)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from None
    from types import SimpleNamespace
    settings = SimpleNamespace(**{key: getattr(config, key) for key in default_values()})
    auth_body = body.get("credenciais") or {}
    if not isinstance(auth_body, dict):
        raise HTTPException(400, "Informe usuário e senha para acessar o SAT")
    credenciais = CredenciaisSAT(auth_body.get("usuario", ""), auth_body.get("senha", ""))
    try:
        credenciais.validar()
    except ValueError:
        raise HTTPException(400, "Informe usuário e senha para acessar o SAT") from None
    job_id = str(uuid.uuid4())[:8]
    job = Job(id=job_id, status="running")
    JOBS[job_id] = job
    handler = JobLogHandler(job_id)

    def progress(index, total, doc, name, stage, results):
        done = [asdict(r) for r in results]
        job.progress = {"index": index, "total": total, "processed": len(done),
                        "current": {"cnpj": doc, "nome": name},
                        "stage": stage, "clients": done}

    def summary(results):
        return {"modo": "consulta", "authentication": job.authentication, "resultados": results, "total": len(results),
                "ok": sum(r["status"] == "ok" for r in results),
                "erro": sum(r["status"] == "erro" for r in results),
                "pendente": sum(r["status"] == "pendente" for r in results)}

    def run_job():
        sat_logger = logging.getLogger("sat")
        sat_logger.addHandler(handler)
        try:
            results = executar_local(params, cancel_event=job.cancel_event,
                                      config=settings, on_progress=progress, credenciais=credenciais,
                                      on_auth=lambda result: setattr(job, "authentication", asdict(result)))
            job.resultado = summary([asdict(r) for r in results])
            job.status = ("cancelado" if job.cancel_event.is_set() else
                          "erro" if job.resultado["erro"] else
                          "pendente" if job.resultado["pendente"] else "ok")
        except ExecucaoCancelada:
            job.resultado = summary(job.progress["clients"])
            job.status = "cancelado"
        except Exception:
            sat_logger.error("Falha na execução SAT; credenciais e detalhes do navegador omitidos")
            job.resultado = summary(job.progress["clients"])
            job.resultado["message"] = "Falha na execução SAT"
            job.status = "erro"
        finally:
            sat_logger.removeHandler(handler)

    job.thread = threading.Thread(target=run_job, daemon=True)
    job.thread.start()
    return {"jobId": job_id}

@app.get("/executar/{job_id}/status")
async def get_job_status(job_id: str):
    """Retorna status de um job."""
    if job_id not in JOBS:
        raise HTTPException(status_code=404, detail="Job não encontrado")

    job = JOBS[job_id]
    logs = list(job.logs)

    resultado = {
        "jobId": job_id,
        "status": job.status,
        "logs": logs,
        "progress": job.progress,
        "authentication": job.authentication,
    }

    if job.resultado:
        resultado["resultado"] = job.resultado

    return resultado

@app.post("/executar/{job_id}/cancelar")
async def cancel_job(job_id: str):
    """Cancela um job em execução."""
    if job_id not in JOBS:
        raise HTTPException(status_code=404, detail="Job não encontrado")

    job = JOBS[job_id]
    if job.status == "running":
        job.cancel_event.set()
        return {"ok": True, "message": "Cancelamento solicitado"}
    return {"ok": False, "message": "Job não está em execução"}

# =============================================================================
# Entry point
# =============================================================================
if __name__ == "__main__":
    try:
        api_port = int(os.environ.get("NFSE_API_PORT", "17432"))
    except ValueError:
        api_port = 17432
    if not 1 <= api_port <= 65535:
        api_port = 17432
    uvicorn.run(app, host="127.0.0.1", port=api_port, log_level="info")
