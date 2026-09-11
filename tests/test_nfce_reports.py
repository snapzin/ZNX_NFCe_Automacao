from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
import pytest

from app.automation.sat_automacao import Parametros
from app.reports import nfce
from app.output_layout import nome_periodo

CNPJ="01234567000189"
NFE="http://www.portalfiscal.inf.br/nfe"


def note(key, value, name="=Produto literal"):
    return f'''<nfeProc xmlns="{NFE}"><NFe><infNFe Id="NFe{key}">
      <ide><mod>65</mod><nNF>1</nNF><serie>001</serie><dhEmi>2026-08-01T10:00:00-03:00</dhEmi></ide>
      <emit><CNPJ>{CNPJ}</CNPJ><xNome>=Empresa literal</xNome></emit>
      <det nItem="1"><prod><xProd>{name}</xProd><qCom>1</qCom><vUnCom>{value}</vUnCom><vProd>{value}</vProd><CFOP>5102</CFOP></prod></det>
      <total><ICMSTot><vNF>{value}</vNF><vProd>{value}</vProd><vPIS>0.10</vPIS><vCOFINS>0.20</vCOFINS></ICMSTot></total>
      </infNFe></NFe><protNFe><infProt><cStat>100</cStat></infProt></protNFe></nfeProc>'''.encode()


def event(key, status="135", day="2026-09-01", kind="110111"):
    return f'''<procEventoNFe xmlns="{NFE}"><evento><infEvento Id="ID{kind}{key}">
      <chNFe>{key}</chNFe><dhEvento>{day}T10:00:00-03:00</dhEvento><tpEvento>{kind}</tpEvento>
      <detEvento><descEvento>Cancelamento</descEvento></detEvento></infEvento></evento>
      <retEvento><infEvento><chNFe>{key}</chNFe><tpEvento>{kind}</tpEvento><cStat>{status}</cStat><nProt>000123</nProt></infEvento></retEvento></procEventoNFe>'''.encode()


@pytest.fixture
def data(tmp_path,monkeypatch):
    # Assinatura é coberta pelos testes criptográficos; aqui isolamos apuração e Excel.
    monkeypatch.setattr(nfce,"preparar_xml",lambda root: nfce.etree.tostring(root))
    folder=tmp_path/CNPJ/"periodo"/"ator_1"
    folder.mkdir(parents=True)
    (folder/"NFCe_2026-08-01_1.xml").write_bytes(note("1"*44,"100.30"))
    (folder/"NFCe_2026-08-01_2.xml").write_bytes(note("2"*44,"50.20"))
    (folder/"Evento_2026-09-01_3.xml").write_bytes(event("1"*44))
    (folder/"Evento_2026-08-02_4.xml").write_bytes(event("2"*44,"136","2026-08-02"))
    return folder


def test_report_matches_decimal_totals_and_cross_period_cancellation(data,tmp_path):
    target=tmp_path/"report.xlsx"
    totals=nfce.gerar_cliente(data,"=Nome literal",CNPJ,"2026-08-01","2026-08-31",{"notas":2,"status":"ok"},target)
    assert totals["valor"]=="150.50"
    assert totals["valor_canceladas"]=="100.30"
    assert totals["valor_sem_cancelamento"]=="50.20"
    assert totals["canceladas"]==1 and totals["eventos_periodo"]==1 and totals["eventos_relacionados"]==2
    wb=load_workbook(target)
    assert wb.sheetnames==["Resumo","NFC-e","Canceladas","Itens","Eventos"]
    notes=wb["NFC-e"]
    assert notes["P5"].value==CNPJ and notes["P5"].data_type=="s"
    assert notes["V5"].value=="1"*44 and notes["V5"].data_type=="s"
    assert notes["Q5"].value=="=Empresa literal" and notes["Q5"].data_type=="s"
    assert wb["Itens"]["F5"].value=="=Produto literal" and wb["Itens"]["F5"].data_type=="s"
    assert notes["E5"].value==100.30 and 'R$' in notes["E5"].number_format
    assert notes.freeze_panes=="C5" and notes.auto_filter.ref=="A4:W6"
    assert notes["D6"].value=="Sem cancelamento identificado"
    wb.close()


def test_each_company_has_only_report_and_xml_folders(data,tmp_path):
    clients=[{"documento":CNPJ,"nome":"=literal"},{"documento":"22222222000100","nome":"Falhou"},
             {"documento":"33333333000100","nome":"Cancelada execução"}]
    params=Parametros("2026-08-01","2026-08-31",clients)
    results=[{"cnpj":CNPJ,"status":"ok","notas":2,"pasta":str(data)},
             {"cnpj":clients[1]["documento"],"status":"erro","message":"Sem vínculo"}]
    report=nfce.gerar_execucao(params,results,tmp_path,"test-job","cancelado")
    assert set(report)=={"clientes","pastas","erros"}
    for client in clients:
        root=Path(report["pastas"][client["documento"]]["empresa"])
        assert {p.name for p in root.iterdir()}=={"08 - Agosto 2026"}
        period=root/"08 - Agosto 2026"
        assert {p.name for p in period.iterdir()}=={"Relatorio","XMLs"}
        assert len(list((period/"Relatorio").glob("*.xlsx")))==1
    assert Path(report["pastas"][CNPJ]["xml"]).is_dir()


def test_report_rejects_mismatched_count_or_bad_values(data,tmp_path):
    with pytest.raises(ValueError,match="Contagem"):
        nfce.gerar_cliente(data,"Teste",CNPJ,"2026-08-01","2026-08-31",{"notas":3},tmp_path/"wrong.xlsx")
    (data/"NFCe_2026-08-01_1.xml").write_bytes(note("1"*44,"NaN"))
    with pytest.raises(ValueError,match="numérico"):
        nfce.read_client(data,CNPJ,"2026-08-01","2026-08-31")


def test_rejected_cancellation_is_not_deducted(data):
    cache=data/"Evento_2026-09-01_3.xml"
    cache.write_bytes(cache.read_bytes().replace(b'<cStat>135',b'<cStat>573'))
    notes,_,_=nfce.read_client(data,CNPJ,"2026-08-01","2026-08-31")
    assert all(n["situacao"]!="Cancelada" for n in notes)


def test_duplicate_company_names_get_separate_folders(tmp_path):
    clients=[{"documento":CNPJ,"nome":"Mesmo Nome"},{"documento":"22222222000100","nome":"mesmo nome"}]
    report=nfce.gerar_execucao(Parametros("2026-08-01","2026-08-31",clients),[],tmp_path,"job","cancelado")
    roots={Path(data["emp