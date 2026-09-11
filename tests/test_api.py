from contextlib import contextmanager
import logging
import threading
import time
from unittest.mock import MagicMock
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from api import server
from app.automation import sat_automacao
from app.automation.sat_downloader import FluxoNaoMapeado
from app.automation.sat_auth import resultado, check_cancel


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(server.config, "XLSX_PATH", str(tmp_path / "clientes.xlsx"))
    monkeypatch.setattr(server, "_API_TOKEN", "test-session-token")
    session = MagicMock()
    session.autenticar.return_value = resultado("SUCCESS")
    session.abrir_consulta_nfce.side_effect = FluxoNaoMapeado()
    @contextmanager
    def open_session(cfg, cert, event, step):
        session.cancel_event = event
        yield session
    monkeypatch.setattr(sat_automacao, "abrir_sessao", open_session)
    server.JOBS.clear()
    with TestClient(server.app) as c:
        c.headers.update({"x-api-token": "test-session-token"})
        c.session = session
        yield c
    for job in server.JOBS.values():
        job.cancel_event.set()
        if job.thread:
            job.thread.join(3)


def customer(**changes):
    return {"documento": "12345678000195", "nome": "Empresa Teste", "ativo": True, **changes}


def wait_job(client, jobid):
    for _ in range(100):
        data = client.get(f"/executar/{jobid}/status").json()
        if data["status"] != "running":
            return data
        time.sleep(.02)
    pytest.fail("Job did not finish")


def test_token_without_license(client):
    assert client.get("/health").status_code == 200
    assert client.get("/health/auth", headers={"x-api-token": "wrong"}).status_code == 401
    assert client.get("/clientes").status_code == 200
    assert client.get("/clientes", headers={"x-api-token": ""}).status_code == 401
    assert client.get("/license/status").status_code == 404


def test_customer_needs_only_name_cnpj_and_saves_two_columns(client):
    assert client.post("/clientes", json={"clientes": [customer(nome="=Nome literal")]}).status_code == 200
    wb = load_workbook(server.config.XLSX_PATH)
    assert [c.value for c in wb.active[1]] == ["NOME", "CNPJ"]
    assert wb.active["A2"].data_type == "s"
    wb.close()
    saved = client.get("/clientes").json()["clientes"][0]
    assert set(saved) == {"nome", "documento", "ativo"}
    saved["ativo"] = False
    client.post("/clientes", json={"clientes": [saved]})
    assert client.post("/executar", json={}).status_code == 400


@pytest.mark.parametrize("headers,row", [
    (["nome", "cnpj"], ["Importada", "12345678000195"]),
    (["CNPJ", "NOME"], ["12345678000195", "Importada"]),
    (["NOME", "CNPJ", "USUARIO_NFSE", "SENHA_NFSE"], ["Importada", "12345678000195", "ignored", "secret"]),
])
def test_import_two_columns_or_legacy_extra_columns(client, tmp_path, headers, row):
    p = tmp_path / "import.xlsx"
    wb = Workbook()
    wb.active.append(headers)
    wb.active.append(row)
    wb.save(p)
    wb.close()
    res = client.post("/clientes/importar", json={"path": str(p)})
    assert res.status_code == 200
    assert res.json()["clientes"] == [customer(nome="Importada")]
    assert "secret" not in res.text


def test_job_authenticates_once_and_reports_download_pending(client, caplog):
    with caplog.at_level(logging.INFO):
        client.post("/clientes", json={"clientes": [customer(), customer(documento="22345678000195")]})
        response = client.post("/executar", json={"credenciais": {"usuario": "test-user", "senha": "test-secret"}})
        assert response.status_code == 200
        data = wait_job(client, response.json()["jobId"])
        assert data["authentication"]["success"]
        assert data["status"] == "pendente"
        assert data["resultado"]["pendente"] == 2 and data["resultado"]["ok"] == 0
        assert data["progress"]["processed"] == 2
        client.session.autenticar.assert_called_once()
        assert "test-secret" not in str(data) + caplog.text


def test_shared_login_failure_does_not_try_client_logins(client):
    client.session.autenticar.return_value = resultado("INVALID_CREDENTIALS")
    client.post("/clientes", json={"clientes": [customer(), customer(documento="22345678000195")]})
    jobid = client.post("/executar", json={"credenciais": {"usuario": "test-user", "senha": "test-secret"}}).json()["jobId"]
    data = wait_job(client, jobid)
    assert data["resultado"]["erro"] == 2
    assert not data["authentication"]["success"]
    client.session.autenticar.assert_called_once()


def test_cancel_shared_login_and_reject_concurrent_job(client):
    entered = threading.Event()
    def authenticate():
        entered.set()
        client.session.cancel_event.wait(3)
        check_cancel(client.session.cancel_event)
    client.session.autenticar.side_effect = authenticate
    client.post("/clientes", json={"clientes": [customer()]})
    jobid = client.post("/executar", json={"credenciais": {"usuario": "test-user", "senha": "test-secret"}}).json()["jobId"]
    assert entered.wait(2)
    assert client.post("/executar", json={}).status_code == 409
    assert client.post(f"/executar/{jobid}/cancelar").json()["ok"]
    assert wait_job(client, jobid)["status"] == "cancelado"


def test_bad_period_or_missing_credentials_rejected_before_job(client):
    client.post("/clientes", json={"clientes": [customer()]})
    assert client.post("/executar", json={"dataInicio": "invalid", "dataFim": "invalid"}).status_code == 400
    assert client.post("/executar", json={"credenciais": {"usuario": "test-user"}}).status_code == 400
    assert client.post("/executar", json={"tipoLogin": "certificado"}).status_code == 400
    assert not server.JOBS


def test_validation_does_not_echo_request_secrets(client, caplog):
    response = client.post("/executar", json=[{"senha": "do-not-echo"}])
    assert response.status_code == 422
    assert "do-not-echo" not in response.text + caplog.text


def test_write_failure_is_not_reported_as_saved(client, monkeypatch, caplog):
    def fail(*a):
        raise OSError("write-error test-secret")
    monkeypatch.setattr(server, "_salvar_clientes_local", fail)
    response = client.post("/clientes", json={"clientes": [customer()]})
    assert response.status_code == 500
    assert "test-secret" not in response.text + caplog.text


def test_login_password_not_in_config(client):
    client.post("/config", json={"SAT_SENHA": "do-not-save", "SAT_USUARIO": "do-not-save"})
    assert "do-not-save" not in server.ensure_config_file().read_text(encoding="utf-8")
    assert "do-not-save" not in client.get("/config").text


def test_selector_config_and_official_url(client):
    selector = 'a[data-label="observado"]'
    response = client.post("/config", json={"SAT_SELECTOR_LOGIN_OK": selector, "SAT_LOGIN_URL": "https://example.invalid/"})
    assert response.status_code == 200
    assert server.config.SAT_SELECTOR_LOGIN_OK == selector
    assert server.config.SAT_LOGIN_URL == "https://sat.sef.sc.gov.br/tax.NET/Login.aspx"
    assert client.post("/config", json={"SAT_SELECTOR_LOGIN_OK": ""}).status_code == 200


@pytest.mark.parametrize("credentials", ["not-a-dict", {"usuario": ["do-not-echo"], "senha": "do-not-echo"}, {"usuario": "user", "senha": {"value": "do-not-echo"}}])
def test_malformed_credentials_rejected_safely(client, credentials):
    client.post("/clientes", json={"clientes": [customer()]})
    response = client.post("/executar", json={"credenciais": credentials})
    assert response.status_code == 400
    assert "do-not-echo" not in response.text
    assert not server.JOBS
